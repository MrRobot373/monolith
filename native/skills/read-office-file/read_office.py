import sys
import os
import urllib.parse

# Reconfigure stdout to use UTF-8 to prevent encoding crashes on Windows/CMD
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass

def print_help():
    print("Usage:")
    print("  python read_office.py info <file_path>                   - Show file structure and metadata")
    print("  python read_office.py read <file_path> [sheet_name/index] [max_rows] - Read Excel sheet or Word document text")
    sys.exit(1)

def handle_excel_info(file_path):
    import pandas as pd
    try:
        xl = pd.ExcelFile(file_path)
        print(f"File Type: Excel Spreadsheet (.xlsx)")
        print(f"Sheets available in workbook:")
        for sheet in xl.sheet_names:
            # Get sheet details
            df = pd.read_excel(xl, sheet_name=sheet)
            print(f"  - Sheet: \"{sheet}\" (Rows: {df.shape[0]}, Columns: {df.shape[1]})")
            print(f"    Columns: {', '.join(map(str, df.columns[:10]))}" + ("..." if len(df.columns) > 10 else ""))
    except Exception as e:
        print(f"Error reading Excel info: {e}")

def handle_excel_read(file_path, sheet_name=None, max_rows=100):
    import pandas as pd
    try:
        xl = pd.ExcelFile(file_path)
        if not sheet_name:
            sheet_name = xl.sheet_names[0]
        elif sheet_name.isdigit():
            idx = int(sheet_name)
            if 0 <= idx < len(xl.sheet_names):
                sheet_name = xl.sheet_names[idx]
            else:
                print(f"Error: Sheet index {idx} out of range (0 to {len(xl.sheet_names)-1}).")
                sys.exit(1)

        print(f"Reading Sheet: \"{sheet_name}\"")
        df = pd.read_excel(xl, sheet_name=sheet_name)
        
        # Limit rows to avoid blowing up context
        total_rows = len(df)
        df_limited = df.head(max_rows)
        
        # Use pandas markdown formatting if tabulate is installed, else simple fallback
        try:
            markdown_table = df_limited.to_markdown(index=False)
            print(markdown_table)
        except ImportError:
            # Fallback to CSV-like table format if tabulate package isn't present
            print(" | ".join(map(str, df_limited.columns)))
            print("-" * (len(df_limited.columns) * 10))
            for _, row in df_limited.iterrows():
                print(" | ".join(map(str, row.values)))
                
        if total_rows > max_rows:
            print(f"\n... (Showing {max_rows} of {total_rows} total rows. Use larger max_rows parameter to see more) ...")
    except Exception as e:
        print(f"Error reading Excel sheet: {e}")

def handle_docx_info(file_path):
    import docx
    try:
        doc = docx.Document(file_path)
        print(f"File Type: Word Document (.docx)")
        print(f"Paragraphs count: {len(doc.paragraphs)}")
        print(f"Tables count: {len(doc.tables)}")
        
        # List major headings if any
        headings = []
        for p in doc.paragraphs:
            if p.style.name.startswith("Heading") and p.text.strip():
                headings.append(f"  - {p.style.name}: \"{p.text.strip()}\"")
        if headings:
            print("Document Outline:")
            print("\n".join(headings[:20]))
            if len(headings) > 20:
                print("  ...")
    except Exception as e:
        print(f"Error reading Word info: {e}")

def handle_docx_read(file_path):
    import docx
    try:
        doc = docx.Document(file_path)
        print(f"Reading Word Document: {os.path.basename(file_path)}\n")
        
        # We can read all paragraphs and tables sequentially
        # Let's extract them in order of appearance
        for element in doc.element.body:
            if element.tag.endswith('p'):
                p = docx.text.paragraph.Paragraph(element, doc)
                text = p.text.strip()
                if text:
                    # Format headings in markdown
                    if p.style.name.startswith("Heading 1"):
                        print(f"\n# {text}\n")
                    elif p.style.name.startswith("Heading 2"):
                        print(f"\n## {text}\n")
                    elif p.style.name.startswith("Heading 3"):
                        print(f"\n### {text}\n")
                    else:
                        print(text)
            elif element.tag.endswith('tbl'):
                tbl = docx.table.Table(element, doc)
                print("\n")
                for row_idx, row in enumerate(tbl.rows):
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    # Dedup consecutive identical cells if merged
                    cleaned_cells = []
                    for cell in cells:
                        if not cleaned_cells or cleaned_cells[-1] != cell:
                            cleaned_cells.append(cell)
                        else:
                            cleaned_cells.append("")
                    print("| " + " | ".join(cleaned_cells) + " |")
                    if row_idx == 0:
                        print("|" + "---|"*len(cells))
                print("\n")
    except Exception as e:
        print(f"Error reading Word document: {e}")

def main():
    if len(sys.argv) < 3:
        print_help()

    cmd = sys.argv[1].lower()
    file_path = sys.argv[2]

    if not os.path.exists(file_path):
        print(f"Error: File '{file_path}' does not exist.")
        sys.exit(1)

    _, ext = os.path.splitext(file_path.lower())

    if cmd == "info":
        if ext == ".xlsx" or ext == ".xls":
            handle_excel_info(file_path)
        elif ext == ".docx":
            handle_docx_info(file_path)
        else:
            print(f"Unsupported file extension: {ext}. Only .xlsx and .docx are supported.")
    elif cmd == "read":
        if ext == ".xlsx" or ext == ".xls":
            sheet = sys.argv[3] if len(sys.argv) > 3 else None
            max_rows = int(sys.argv[4]) if len(sys.argv) > 4 else 100
            handle_excel_read(file_path, sheet, max_rows)
        elif ext == ".docx":
            handle_docx_read(file_path)
        else:
            print(f"Unsupported file extension: {ext}. Only .xlsx and .docx are supported.")
    else:
        print_help()

if __name__ == '__main__':
    main()
