import sys
import os
import json
import urllib.parse

# Reconfigure stdout to use UTF-8 to prevent encoding crashes on Windows/CMD
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass

def print_help():
    print("Office Document Manager Utility")
    print("Usage:")
    print("  python office_manager.py info <file_path>")
    print("  python office_manager.py read <file_path> [extra_arg1] [extra_arg2]")
    print("  python office_manager.py create <file_path> <json_data_or_file>")
    print("  python office_manager.py edit <file_path> <json_edit_data_or_file>")
    print("\nSupported formats: Excel (.xlsx), Word (.docx), PDF (.pdf), PowerPoint (.pptx)")
    sys.exit(1)

# ==============================================================================
# EXCEL (XLSX) HANDLING
# ==============================================================================
def handle_excel_info(file_path):
    import pandas as pd
    xl = pd.ExcelFile(file_path)
    print(f"Format: Excel (.xlsx)")
    print(f"Worksheets:")
    for sheet in xl.sheet_names:
        df = pd.read_excel(xl, sheet_name=sheet)
        print(f"  - \"{sheet}\": {df.shape[0]} rows x {df.shape[1]} columns")
        print(f"    Columns: {', '.join(map(str, df.columns[:8]))}" + ("..." if len(df.columns) > 8 else ""))

def handle_excel_read(file_path, sheet_name=None, max_rows=100):
    import pandas as pd
    xl = pd.ExcelFile(file_path)
    if not sheet_name:
        sheet_name = xl.sheet_names[0]
    elif sheet_name.isdigit():
        idx = int(sheet_name)
        sheet_name = xl.sheet_names[idx]
        
    df = pd.read_excel(xl, sheet_name=sheet_name)
    print(f"Sheet: \"{sheet_name}\" (Total rows: {len(df)})")
    
    df_limited = df.head(int(max_rows))
    try:
        print(df_limited.to_markdown(index=False))
    except ImportError:
        print(" | ".join(map(str, df_limited.columns)))
        print("-" * (len(df_limited.columns) * 10))
        for _, row in df_limited.iterrows():
            print(" | ".join(map(str, row.values)))
            
    if len(df) > int(max_rows):
        print(f"\n... Showing top {max_rows} rows of {len(df)} total rows. ...")

def handle_excel_create(file_path, json_data):
    import pandas as pd
    # json_data can be a list of dicts or a dict mapping sheet name -> list of dicts
    if isinstance(json_data, list):
        df = pd.DataFrame(json_data)
        df.to_excel(file_path, index=False)
        print(f"Created Excel file '{file_path}' with {len(df)} rows.")
    elif isinstance(json_data, dict):
        with pd.ExcelWriter(file_path) as writer:
            for sheet, rows in json_data.items():
                df = pd.DataFrame(rows)
                df.to_excel(writer, sheet_name=sheet, index=False)
                print(f"Created sheet '{sheet}' with {len(df)} rows.")
        print(f"Created Excel workbook '{file_path}' successfully.")

def handle_excel_edit(file_path, edit_data):
    import pandas as pd
    xl = pd.ExcelFile(file_path)
    sheets_data = {}
    for sheet in xl.sheet_names:
        sheets_data[sheet] = pd.read_excel(xl, sheet_name=sheet)

    # edit_data can contain actions: "update_cell" or "append_rows"
    action = edit_data.get("action", "")
    sheet = edit_data.get("sheet", xl.sheet_names[0])
    
    if sheet not in sheets_data:
        print(f"Error: Sheet '{sheet}' not found.")
        sys.exit(1)
        
    df = sheets_data[sheet]

    if action == "update_cell":
        row = int(edit_data["row"])
        col = edit_data["col"] # can be index or name
        val = edit_data["val"]
        if isinstance(col, int):
            df.iat[row, col] = val
        else:
            df.at[row, col] = val
        print(f"Updated cell {row}, {col} with value: {val}")
    elif action == "append_rows":
        rows = edit_data["rows"]
        new_df = pd.DataFrame(rows)
        df = pd.concat([df, new_df], ignore_index=True)
        sheets_data[sheet] = df
        print(f"Appended {len(rows)} rows to sheet '{sheet}'.")
    else:
        print(f"Unknown Excel edit action: {action}")
        sys.exit(1)

    with pd.ExcelWriter(file_path) as writer:
        for s, d in sheets_data.items():
            d.to_excel(writer, sheet_name=s, index=False)
    print(f"Saved edits to Excel file '{file_path}'.")

# ==============================================================================
# WORD (DOCX) HANDLING
# ==============================================================================
def handle_docx_info(file_path):
    import docx
    doc = docx.Document(file_path)
    print(f"Format: Word Document (.docx)")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    if headings:
        print("Structure Headings:")
        for h in headings[:15]:
            print(f"  - {h}")

def handle_docx_read(file_path):
    import docx
    doc = docx.Document(file_path)
    for element in doc.element.body:
        if element.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(element, doc)
            text = p.text.strip()
            if text:
                if p.style.name.startswith("Heading 1"):
                    print(f"\n# {text}\n")
                elif p.style.name.startswith("Heading 2"):
                    print(f"\n## {text}\n")
                elif p.style.name.startswith("Heading 3"):
                    print(f"\n### {text}\n")
                else:
                    print(text)
        elif element.tag.endswith('tbl'):
            tbl = docx.table.Table(element, doc)
            print("\n")
            for r_idx, row in enumerate(tbl.rows):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                # Filter consec merges
                c_cells = []
                for c in cells:
                    if not c_cells or c_cells[-1] != c:
                        c_cells.append(c)
                    else:
                        c_cells.append("")
                print("| " + " | ".join(c_cells) + " |")
                if r_idx == 0:
                    print("|" + "---|"*len(cells))
            print("\n")

def handle_docx_create(file_path, json_data):
    import docx
    doc = docx.Document()
    # json_data is a list of elements: [{"type": "heading", "level": 1, "text": "..."}, {"type": "paragraph", "text": "..."}]
    for el in json_data:
        t = el.get("type", "paragraph").lower()
        if t == "heading":
            doc.add_heading(el.get("text", ""), level=int(el.get("level", 1)))
        elif t == "paragraph":
            doc.add_paragraph(el.get("text", ""))
        elif t == "bullet":
            doc.add_paragraph(el.get("text", ""), style='List Bullet')
        elif t == "table":
            headers = el.get("headers", [])
            rows = el.get("rows", [])
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for idx, name in enumerate(headers):
                hdr_cells[idx].text = str(name)
            for r_data in rows:
                row_cells = table.add_row().cells
                for idx, val in enumerate(r_data):
                    row_cells[idx].text = str(val)
    doc.save(file_path)
    print(f"Created Word document '{file_path}' successfully.")

def handle_docx_edit(file_path, edit_data):
    import docx
    doc = docx.Document(file_path)
    action = edit_data.get("action", "")
    if action == "append":
        for el in edit_data.get("elements", []):
            t = el.get("type", "paragraph").lower()
            if t == "heading":
                doc.add_heading(el.get("text", ""), level=int(el.get("level", 1)))
            elif t == "paragraph":
                doc.add_paragraph(el.get("text", ""))
            elif t == "bullet":
                doc.add_paragraph(el.get("text", ""), style='List Bullet')
    elif action == "replace":
        old_text = edit_data["old_text"]
        new_text = edit_data["new_text"]
        count = 0
        for p in doc.paragraphs:
            if old_text in p.text:
                p.text = p.text.replace(old_text, new_text)
                count += 1
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if old_text in cell.text:
                        cell.text = cell.text.replace(old_text, new_text)
                        count += 1
        print(f"Replaced text in {count} locations.")
    else:
        print(f"Unknown Word edit action: {action}")
        sys.exit(1)
    doc.save(file_path)
    print(f"Saved edits to Word document '{file_path}'.")

# ==============================================================================
# PDF HANDLING
# ==============================================================================
def handle_pdf_info(file_path):
    import pypdf
    reader = pypdf.PdfReader(file_path)
    print(f"Format: PDF (.pdf)")
    print(f"Pages: {len(reader.pages)}")
    meta = reader.metadata
    if meta:
        print("Metadata:")
        for k, v in meta.items():
            print(f"  - {k}: {v}")

def handle_pdf_read(file_path):
    import pdfplumber
    with pdfplumber.open(file_path) as pdf:
        print(f"Reading PDF: {os.path.basename(file_path)} ({len(pdf.pages)} pages)\n")
        for idx, page in enumerate(pdf.pages, 1):
            print(f"--- Page {idx} ---")
            text = page.extract_text()
            if text:
                print(text)
            tables = page.extract_tables()
            if tables:
                print("\n[Tables Found on Page]")
                for table in tables:
                    for r in table:
                        cleaned_row = [str(cell or "").strip().replace("\n", " ") for cell in r]
                        print("| " + " | ".join(cleaned_row) + " |")
                    print()
            print()

def handle_pdf_create(file_path, json_data):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    title = json_data.get("title", "Document")
    elements_data = json_data.get("elements", [])

    doc = SimpleDocTemplate(file_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1A365D'),
        spaceAfter=15
    )
    
    normal_style = styles['Normal']
    normal_style.fontSize = 11
    normal_style.leading = 15
    normal_style.spaceAfter = 8

    story = []
    story.append(Paragraph(title, heading_style))
    story.append(Spacer(1, 10))

    for el in elements_data:
        t = el.get("type", "paragraph").lower()
        if t == "heading":
            level = int(el.get("level", 1))
            size = 18 - (level * 2)
            style = ParagraphStyle(
                f'H_{level}',
                parent=styles['Heading2'],
                fontSize=size,
                leading=size + 4,
                textColor=colors.HexColor('#2B6CB0'),
                spaceBefore=12,
                spaceAfter=6
            )
            story.append(Paragraph(el.get("text", ""), style))
        elif t == "paragraph":
            story.append(Paragraph(el.get("text", ""), normal_style))
        elif t == "table":
            headers = el.get("headers", [])
            rows = el.get("rows", [])
            table_data = [headers] + rows
            t_el = Table(table_data)
            t_el.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#4A5568')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('BOTTOMPADDING', (0,0), (-1,0), 6),
                ('TOPPADDING', (0,0), (-1,0), 6),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#F7FAFC'), colors.white]),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 9),
            ]))
            story.append(Spacer(1, 5))
            story.append(t_el)
            story.append(Spacer(1, 5))

    doc.build(story)
    print(f"Successfully generated PDF: {file_path}")

# ==============================================================================
# POWERPOINT (PPTX) HANDLING
# ==============================================================================
def handle_pptx_info(file_path):
    import pptx
    prs = pptx.Presentation(file_path)
    print(f"Format: PowerPoint Presentation (.pptx)")
    print(f"Slides: {len(prs.slides)}")
    print("Slides Outline:")
    for idx, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text if slide.shapes.title else "Untitled Slide"
        print(f"  Slide {idx}: \"{title}\"")

def handle_pptx_read(file_path):
    import pptx
    prs = pptx.Presentation(file_path)
    print(f"Reading Presentation: {os.path.basename(file_path)} ({len(prs.slides)} slides)\n")
    for idx, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text if slide.shapes.title else "Untitled Slide"
        print(f"--- Slide {idx}: {title} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        indent = paragraph.level
                        print("  " * indent + f"- {text}")
        print()

def handle_pptx_create(file_path, json_data):
    import pptx
    prs = pptx.Presentation()
    # json_data is a list of slides: [{"title": "...", "bullets": ["...", "..."]}]
    for s_data in json_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s_data.get("title", "")
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        bullets = s_data.get("bullets", [])
        for i, b in enumerate(bullets):
            if i == 0:
                p = tf.paragraphs[0]
                p.text = b
            else:
                p = tf.add_paragraph()
                p.text = b
    prs.save(file_path)
    print(f"Created PowerPoint presentation '{file_path}' successfully.")

def handle_pptx_edit(file_path, edit_data):
    import pptx
    prs = pptx.Presentation(file_path)
    action = edit_data.get("action", "")
    
    if action == "add_slide":
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = edit_data.get("title", "")
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        bullets = edit_data.get("bullets", [])
        for i, b in enumerate(bullets):
            if i == 0:
                p = tf.paragraphs[0]
                p.text = b
            else:
                p = tf.add_paragraph()
                p.text = b
        print(f"Added slide: '{edit_data.get('title')}'")
    else:
        print(f"Unknown PPTX edit action: {action}")
        sys.exit(1)
        
    prs.save(file_path)
    print(f"Saved edits to PowerPoint file '{file_path}'.")

# ==============================================================================
# MAIN ROUTING
# ==============================================================================
def main():
    if len(sys.argv) < 3:
        print_help()

    cmd = sys.argv[1].lower()
    file_path = sys.argv[2]

    file_path = os.path.abspath(file_path)

    json_data = None
    if cmd in ["create", "edit"]:
        if len(sys.argv) < 4:
            print(f"Error: JSON data or file path required for command '{cmd}'.")
            sys.exit(1)
        data_arg = sys.argv[3]
        if os.path.exists(data_arg):
            with open(data_arg, 'r', encoding='utf-8') as f:
                json_data = json.load(f)
        else:
            try:
                json_data = json.loads(data_arg)
            except Exception as e:
                print(f"Error parsing JSON data: {e}")
                sys.exit(1)

    _, ext = os.path.splitext(file_path.lower())

    if cmd == "info":
        if ext == ".xlsx" or ext == ".xls":
            handle_excel_info(file_path)
        elif ext == ".docx":
            handle_docx_info(file_path)
        elif ext == ".pdf":
            handle_pdf_info(file_path)
        elif ext == ".pptx":
            handle_pptx_info(file_path)
        else:
            print(f"Unsupported format: {ext}")
            sys.exit(1)
            
    elif cmd == "read":
        if ext == ".xlsx" or ext == ".xls":
            sheet = sys.argv[3] if len(sys.argv) > 3 else None
            max_rows = int(sys.argv[4]) if len(sys.argv) > 4 else 100
            handle_excel_read(file_path, sheet, max_rows)
        elif ext == ".docx":
            handle_docx_read(file_path)
        elif ext == ".pdf":
            handle_pdf_read(file_path)
        elif ext == ".pptx":
            handle_pptx_read(file_path)
        else:
            print(f"Unsupported format: {ext}")
            sys.exit(1)

    elif cmd == "create":
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        if ext == ".xlsx" or ext == ".xls":
            handle_excel_create(file_path, json_data)
        elif ext == ".docx":
            handle_docx_create(file_path, json_data)
        elif ext == ".pdf":
            handle_pdf_create(file_path, json_data)
        elif ext == ".pptx":
            handle_pptx_create(file_path, json_data)
        else:
            print(f"Unsupported format: {ext}")
            sys.exit(1)

    elif cmd == "edit":
        if ext == ".xlsx" or ext == ".xls":
            handle_excel_edit(file_path, json_data)
        elif ext == ".docx":
            handle_docx_edit(file_path, json_data)
        elif ext == ".pptx":
            handle_pptx_edit(file_path, json_data)
        else:
            print(f"Unsupported format for editing: {ext}")
            sys.exit(1)
    else:
        print_help()

if __name__ == '__main__':
    main()
