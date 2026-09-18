#!/usr/bin/env python3
"""Create, edit, inspect and validate Excel (.xlsx) workbooks.

    python xlsx_tool.py create   - out.xlsx        < spec.json
    python xlsx_tool.py edit     in.xlsx out.xlsx - < ops.json
    python xlsx_tool.py inspect  in.xlsx [sheet]
    python xlsx_tool.py validate in.xlsx

`-` reads the JSON from stdin, so a workbook is one command with no
intermediate file.

FORMULAS ARE KEPT AS FORMULAS. Write "=SUM(B2:B10)" and Excel evaluates it on
open; this tool never pre-computes a value into a cell, because that silently
turns a live model into stale text. `inspect` shows formulas by default and
takes `--values` to show the last-cached results instead.
"""

import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("error: openpyxl is not installed. Run: python -m pip install openpyxl")

HEADER_FILL = PatternFill("solid", fgColor="DDE5E0")


def fail(message):
    sys.exit(f"error: {message}")


def read_json(source):
    raw = sys.stdin.read() if source == "-" else None
    if raw is None:
        try:
            raw = Path(source).read_text(encoding="utf-8")
        except OSError as exc:
            fail(f"cannot read {source}: {exc}")
    if not raw.strip():
        fail("no JSON received; pipe it in on stdin or pass a file path")
    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        fail(f"input is not valid JSON: {exc}")


def open_wb(path, data_only=False):
    try:
        return load_workbook(path, data_only=data_only)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises varied exceptions
        fail(f"cannot open {path} as .xlsx: {exc}")


def autosize(sheet):
    """Widen columns to their content — unreadable ##### columns are a bug."""
    for column in sheet.columns:
        longest = 0
        letter = None
        for cell in column:
            if letter is None:
                letter = cell.column_letter
            value = cell.value
            if value is not None:
                longest = max(longest, len(str(value)))
        if letter:
            sheet.column_dimensions[letter].width = min(60, max(10, longest + 2))


def write_sheet(sheet, spec, index):
    headers = spec.get("headers") or []
    rows = spec.get("rows") or []
    if headers:
        sheet.append([str(h) for h in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
    for row in rows:
        if not isinstance(row, list):
            fail(f'sheets[{index}].rows entries must be lists, got {type(row).__name__}')
        sheet.append(row)
    for ref, formula in (spec.get("formulas") or {}).items():
        sheet[ref] = formula
    if headers and rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    autosize(sheet)


# ---------------------------------------------------------------- create ----

def cmd_create(args):
    if len(args) != 2:
        fail("usage: create <spec.json|-> <out.xlsx>")
    spec, target = read_json(args[0]), args[1]
    sheets = spec.get("sheets") if isinstance(spec, dict) else None
    if not isinstance(sheets, list) or not sheets:
        fail('"sheets" must be a non-empty list')

    book = Workbook()
    book.remove(book.active)
    for index, sheet_spec in enumerate(sheets):
        if not isinstance(sheet_spec, dict):
            fail(f"sheets[{index}] must be an object")
        name = str(sheet_spec.get("name") or f"Sheet{index + 1}")[:31]
        write_sheet(book.create_sheet(title=name), sheet_spec, index)

    save(book, target, f"{len(sheets)} sheet(s)")


# ------------------------------------------------------------------ edit ----

def cmd_edit(args):
    if len(args) != 3:
        fail("usage: edit <in.xlsx> <out.xlsx> <ops.json|->")
    source, target = args[0], args[1]
    payload = read_json(args[2])
    ops = payload.get("ops") if isinstance(payload, dict) else payload
    if not isinstance(ops, list) or not ops:
        fail('"ops" must be a non-empty list')

    book = open_wb(source)
    applied = []
    for index, op in enumerate(ops):
        if not isinstance(op, dict):
            fail(f"ops[{index}] must be an object")
        kind = op.get("op")

        if kind == "set_cell":
            sheet = pick_sheet(book, op, index)
            ref = op.get("ref")
            if not ref:
                fail(f'ops[{index}] "set_cell" needs "ref" (e.g. "B4")')
            sheet[ref] = op.get("value")
            applied.append(f"set {sheet.title}!{ref}")
        elif kind == "append_row":
            sheet = pick_sheet(book, op, index)
            row = op.get("values")
            if not isinstance(row, list):
                fail(f'ops[{index}] "append_row" needs "values" as a list')
            sheet.append(row)
            applied.append(f"append {sheet.title}")
        elif kind == "add_sheet":
            name = str(op.get("name") or "Sheet")[:31]
            write_sheet(book.create_sheet(title=name), op, index)
            applied.append(f"add sheet {name}")
        elif kind == "delete_sheet":
            sheet = pick_sheet(book, op, index)
            book.remove(sheet)
            applied.append(f"delete sheet {sheet.title}")
        elif kind == "rename_sheet":
            sheet = pick_sheet(book, op, index)
            sheet.title = str(op.get("to") or sheet.title)[:31]
            applied.append(f"rename -> {sheet.title}")
        else:
            fail(f"ops[{index}]: unknown op {kind!r}. Supported: set_cell, "
                 "append_row, add_sheet, delete_sheet, rename_sheet")

    save(book, target, "; ".join(applied))


def pick_sheet(book, op, index):
    name = op.get("sheet")
    if name is None:
        return book.worksheets[0]
    if name not in book.sheetnames:
        fail(f"ops[{index}]: no sheet named {name!r}. Present: {', '.join(book.sheetnames)}")
    return book[name]


# --------------------------------------------------------------- inspect ----

def cmd_inspect(args):
    values = "--values" in args
    args = [a for a in args if a != "--values"]
    if not args:
        fail("usage: inspect <in.xlsx> [sheet] [--values]")
    book = open_wb(args[0], data_only=values)
    wanted = args[1] if len(args) > 1 else None

    print(f"sheets: {', '.join(book.sheetnames)}")
    for sheet in book.worksheets:
        if wanted and sheet.title != wanted:
            continue
        print(f"--- {sheet.title}: {sheet.max_row} rows x {sheet.max_column} cols")
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            print("   " + " | ".join(c[:22] for c in cells))
        if sheet.max_row > 12:
            print(f"   … {sheet.max_row - 12} more row(s)")


# -------------------------------------------------------------- validate ----

def cmd_validate(args):
    if len(args) != 1:
        fail("usage: validate <in.xlsx>")
    path = Path(args[0])
    if not path.exists():
        fail(f"{path} does not exist")

    import zipfile
    if not zipfile.is_zipfile(path):
        fail(f"{path} is not a zip container, so it is not a real .xlsx "
             "(a .csv or .xls renamed to .xlsx will do this)")
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
        if broken:
            fail(f"corrupt entry in archive: {broken}")
        if "xl/workbook.xml" not in archive.namelist():
            fail("missing xl/workbook.xml — the file is not a valid workbook")

    book = open_wb(path)
    issues = []
    for sheet in book.worksheets:
        if sheet.max_row == 1 and sheet.max_column == 1 and sheet["A1"].value is None:
            issues.append(f"sheet {sheet.title!r} is empty")
    print(f"valid: {path.resolve()}")
    print(f"  sheets: {len(book.worksheets)} ({', '.join(book.sheetnames)})")
    for issue in issues:
        print(f"  warning: {issue}")


# ------------------------------------------------------------------------ --

def save(book, target, summary):
    out = Path(target).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        book.save(out)
    except OSError as exc:
        fail(f"cannot write {out}: {exc}")
    print(f"wrote {out} ({summary})")


COMMANDS = {
    "create": cmd_create,
    "edit": cmd_edit,
    "inspect": cmd_inspect,
    "validate": cmd_validate,
}


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in COMMANDS:
        sys.exit(f"usage: xlsx_tool.py <{'|'.join(COMMANDS)}> ...")
    COMMANDS[sys.argv[1]](sys.argv[2:])


if __name__ == "__main__":
    main()
